from __future__ import annotations

import csv
import io
import json
from pathlib import Path


class UnsupportedDocumentError(ValueError):
    pass


TEXT_SUFFIXES = {".txt", ".md", ".log"}


def extract_document_text(filename: str, data: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return data.decode("utf-8-sig").strip()
    if suffix == ".csv":
        rows = csv.reader(io.StringIO(data.decode("utf-8-sig")))
        return "\n".join(" | ".join(cell.strip() for cell in row) for row in rows).strip()
    if suffix == ".json":
        return json.dumps(json.loads(data.decode("utf-8-sig")), ensure_ascii=False, indent=2)
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n\n".join(page for page in pages if page)
        if not text:
            raise UnsupportedDocumentError("No extractable PDF text found; scanned PDFs require OCR.")
        return text
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        blocks = []
        for sheet in workbook.worksheets:
            rows = [" | ".join("" if v is None else str(v) for v in row)
                    for row in sheet.iter_rows(values_only=True)]
            blocks.append(f"[Sheet: {sheet.title}]\n" + "\n".join(r for r in rows if r.strip()))
        return "\n\n".join(blocks).strip()
    raise UnsupportedDocumentError("Unsupported document type.")
